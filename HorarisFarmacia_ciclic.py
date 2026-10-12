import sys
import os
if sys.stdout is None:
    sys.stdout = open(os.devnull, "w")
if sys.stderr is None:
    sys.stderr = open(os.devnull, "w")
if sys.stdin is None:
    sys.stdin = open(os.devnull, "r")
    
    
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import pulp
import json
import os
import random

from PIL import Image, ImageDraw, ImageFont, ImageTk

import platform

_CACHE_ICONES = {}


def _carregar_font_emoji(mida):
    """Return (font, font_size) for the current OS, or (None, 0)."""
    sistema = platform.system()
    if sistema == "Windows":
        candidats = [("seguiemj.ttf", mida)]
    elif sistema == "Darwin":
        # Apple Color Emoji is a bitmap font: only size 160 is valid.
        candidats = [("/System/Library/Fonts/Apple Color Emoji.ttc", 160)]
    else:
        # NotoColorEmoji is also a bitmap font (size 109).
        candidats = [
            ("/usr/share/fonts/truetype/noto/NotoColorEmoji.ttf", 109),
            ("/usr/share/fonts/noto/NotoColorEmoji.ttf", 109),
        ]
    for ruta, mida_font in candidats:
        try:
            return ImageFont.truetype(ruta, mida_font), mida_font
        except Exception:
            continue
    return None, 0


def icona_emoji(emoji_char, mida=18):
    """Render an emoji as a PhotoImage. Never returns None."""
    clau = (emoji_char, mida)
    if clau in _CACHE_ICONES:
        return _CACHE_ICONES[clau]

    costat = mida + 6
    img = None
    font, mida_font = _carregar_font_emoji(mida)
    if font is not None:
        try:
            gran = Image.new("RGBA", (mida_font + 40, mida_font + 40), (255, 255, 255, 0))
            ImageDraw.Draw(gran).text((20, 20), emoji_char, font=font, embedded_color=True)
            bbox = gran.getbbox()
            if bbox:
                gran = gran.crop(bbox)
                gran.thumbnail((mida, mida), Image.LANCZOS)
                img = Image.new("RGBA", (costat, costat), (255, 255, 255, 0))
                img.paste(gran, ((costat - gran.width) // 2, (costat - gran.height) // 2), gran)
        except Exception:
            img = None

    try:
        foto = ImageTk.PhotoImage(img) if img is not None else tk.PhotoImage(width=costat, height=costat)
    except Exception:
        foto = tk.PhotoImage(width=costat, height=costat)

    _CACHE_ICONES[clau] = foto
    return foto

try:
    from reportlab.lib.pagesizes import landscape, A3
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    REPORTLAB_OK = True
except Exception:
    REPORTLAB_OK = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except Exception:
    OPENPYXL_OK = False

FITXER_DADES = "dades_farmacia.json"

OPCIONS_PATRO = [
    "-",
    "Preferentment Matí",
    "Només Matí",
    "Preferentment Tarda",
    "Només Tarda",
    "Tot el dia",
    "No ve",
    "Només aquest dia"
]

DIES_CAP_SETMANA = ("Dissabte", "Diumenge")

SOROLL_ALEATORI_MIN = 0.01
SOROLL_ALEATORI_MAX = 0.20
PES_MANCA_LLICENCIAT = 5000
PES_BONUS_ROTACIO = 6
DIFERENCIA_MAX_EQUITAT_ROTACIO = 1

COLOR_ZEBRA_PAR = "#ffffff"
COLOR_ZEBRA_IMPAR = "#eef2f7"


def to_hhmm(h_float):
    h = int(h_float)
    m = int(round((h_float - h) * 60))
    return f"{h:02d}:{m:02d}"


class AppFarmaciaCiclica:
    def __init__(self, root):
        self.root = root
        self.root.title("Gestor d'Horaris Farmàcia - Horari Cíclic")
        self.root.geometry("1470x1040")
        self.root.configure(bg="#f4f6f9")

        self.root.option_add('*Entry.insertBackground', '#0f172a')
        self.root.option_add('*Entry.insertWidth', 2)
        self.root.option_add('*Spinbox.insertBackground', '#0f172a')
        self.root.option_add('*Spinbox.insertWidth', 2)
        self.root.option_add('*Combobox.insertBackground', '#0f172a')
        self.root.option_add('*Combobox.insertWidth', 2)

        self.mida_font = tk.IntVar(value=10)
        self.style = ttk.Style()
        self.style.theme_use('clam')

        self.treballadors = {}
        self.dies = ['Dilluns', 'Dimarts', 'Dimecres', 'Dijous', 'Divendres', 'Dissabte', 'Diumenge']

        self.min_personal_setmana_var = tk.IntVar(value=2)
        self.max_personal_setmana_var = tk.IntVar(value=4)

        self.min_personal_cap_setmana_var = tk.IntVar(value=1)
        self.max_personal_cap_setmana_var = tk.IntVar(value=3)

        self.min_llicenciats_var = tk.IntVar(value=1)
        self.min_llicenciats_cap_setmana_var = tk.IntVar(value=1)

        # ÚNIC paràmetre de rotació: durada del cicle (X setmanes) = freqüència
        # màxima (1 cap de setmana cada X setmanes per treballador de rotació).
        self.num_setmanes_cicle_var = tk.IntVar(value=2)

        self.inici_canvi_var = tk.StringVar(value="13:30")
        self.fi_canvi_var = tk.StringVar(value="14:30")

        self.durada_min_torn_var = tk.DoubleVar(value=4.0)
        self.durada_max_torn_var = tk.DoubleVar(value=8.0)

        self.horaris_ui = {}
        self.combos_patro = {}

        self.res_quadrant_cicle = None
        self.res_resum_cicle = None
        self.torns_per_dia_nom = {}
        self.num_setmanes_generat = 2
        self.avisos_cap_setmana = []
        self.dies_sense_llicenciat = set()
        self.icones = []

        self.crear_barra_superior()

        self.notebook = ttk.Notebook(root)
        self.notebook.pack(fill='both', expand=True, padx=15, pady=(5, 15))

        self.tab_horaris = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_horaris, text=" 1. Apertura i Canvi de Torn ")
        self.crear_tab_horaris()

        self.tab_treb = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_treb, text=" 2. Equip, Patró Setmanal i Rotació ")
        self.crear_tab_treb()

        self.tab_resultats = ttk.Frame(self.notebook)
        self.notebook.add(self.tab_resultats, text=" 3. Horari Cíclic Optimitzat ")
        self.crear_tab_resultats()

        self.aplicar_estils_font()
        self.carregar_dades()
        self.root.protocol("WM_DELETE_WINDOW", self.on_closing)

    # ----------------------------------------------------------------------
    # ZEBRA (files alternades blanc / gris clar) — ajuda visual a taules
    # ----------------------------------------------------------------------
    def _configurar_zebra(self, tree):
        tree.tag_configure('par', background=COLOR_ZEBRA_PAR)
        tree.tag_configure('impar', background=COLOR_ZEBRA_IMPAR)

    def _aplicar_zebra(self, tree):
        for idx, item in enumerate(tree.get_children()):
            tree.item(item, tags=('par' if idx % 2 == 0 else 'impar',))

    COLORS_TORN_UI = {
        "Matí": ("#fef9c3", "#78350f"),
        "Tarda": ("#fed7aa", "#7c2d12"),
        "Matí + Tarda": ("#bfdbfe", "#1e3a8a"),
        "Festa": ("#dcfce7", "#14532d"),
        "-": ("#f5f5f5", "#64748b"),
        "header": ("#1e293b", "#ffffff"),
    }

    def _cella_horari(self, parent, text, row, col, tipus="-", bold=False, avis=False, ample=14):
        bg, fg = self.COLORS_TORN_UI.get(tipus, ("#ffffff", "#334155"))
        if avis:
            bg, fg = "#fee2e2", "#b91c1c"
        lbl = tk.Label(
            parent, text=text, bg=bg, fg=fg, width=ample,
            font=("Segoe UI", self.mida_font.get(), "bold" if bold else "normal"),
            relief="solid", borderwidth=1, padx=4, pady=6
        )
        lbl.grid(row=row, column=col, sticky="nsew")

    # ----------------------------------------------------------------------
    # BARRA SUPERIOR
    # ----------------------------------------------------------------------
    def crear_barra_superior(self):
        f_top = tk.Frame(self.root, bg="#1e293b")
        f_top.pack(fill="x", side="top", ipady=6)

        f_cicle = tk.Frame(f_top, bg="#1e293b")
        f_cicle.pack(side="left", padx=10)

        ic_cicle = icona_emoji("🔁", 16); self.icones.append(ic_cicle)
        tk.Label(f_cicle, image=ic_cicle, text=" Cicle / Rotació Caps de Setmana (X setmanes):", compound="left", font=("Segoe UI", 10, "bold"), fg="#ffffff", bg="#1e293b").pack(side="left", padx=3)
        sp_cicle = ttk.Spinbox(f_cicle, from_=1, to=8, textvariable=self.num_setmanes_cicle_var, width=4,
                                command=self.actualitzar_info_rotacio)
        sp_cicle.pack(side="left", padx=3)
        sp_cicle.bind("<KeyRelease>", lambda e: self.actualitzar_info_rotacio())

        tk.Label(f_cicle, text="(Cicle fix, es repeteix indefinidament, sense mes ni festius)",
                 font=("Segoe UI", 9, "italic"), fg="#94a3b8", bg="#1e293b").pack(side="left", padx=(10, 0))

        f_zoom = tk.Frame(f_top, bg="#1e293b")
        f_zoom.pack(side="right", padx=10)

        ic_zoom = icona_emoji("🔍", 16); self.icones.append(ic_zoom)
        tk.Label(f_zoom, image=ic_zoom, text=" Mida Text:", compound="left", font=("Segoe UI", 10, "bold"), fg="#ffffff", bg="#1e293b").pack(side="left", padx=5)

        btn_minus = tk.Label(f_zoom, text=" − ", font=("Segoe UI", 11, "bold"), fg="#ffffff", bg="#334155", cursor="hand2", padx=6, pady=2)
        btn_minus.pack(side="left", padx=3)
        btn_minus.bind("<Button-1>", lambda e: self.canviar_mida_font(-1))

        self.lbl_font_mida = tk.Label(f_zoom, textvariable=self.mida_font, font=("Segoe UI", 10, "bold"), fg="#38bdf8", bg="#1e293b", width=3)
        self.lbl_font_mida.pack(side="left")

        btn_plus = tk.Label(f_zoom, text=" + ", font=("Segoe UI", 11, "bold"), fg="#ffffff", bg="#334155", cursor="hand2", padx=6, pady=2)
        btn_plus.pack(side="left", padx=3)
        btn_plus.bind("<Button-1>", lambda e: self.canviar_mida_font(1))

    def actualitzar_info_rotacio(self):
        if hasattr(self, 'lbl_rotacio_info'):
            self.lbl_rotacio_info.config(
                text=f"🔁 Rotació: 1 cap de setmana cada {self.num_setmanes_cicle_var.get()} setmanes (paràmetre de dalt)."
            )

    def canviar_mida_font(self, delta):
        nova = max(8, min(16, self.mida_font.get() + delta))
        self.mida_font.set(nova)
        self.aplicar_estils_font()

    def aplicar_estils_font(self):
        f = self.mida_font.get()
        PRIMARY = "#1e293b"
        ACCENT = "#2563eb"
        BG_COLOR = "#ffffff"
        LIGHT_BG = "#f8fafc"
        BORDER_COLOR = "#e2e8f0"

        self.style.configure("TFrame", background=BG_COLOR)
        self.style.configure("TLabelframe", background=BG_COLOR, foreground=PRIMARY)
        self.style.configure("TLabelframe.Label", font=("Segoe UI", f, "bold"), background=BG_COLOR, foreground=PRIMARY)
        self.style.configure("TLabel", background=BG_COLOR, font=("Segoe UI", f), foreground="#334155")
        self.style.configure("TButton", font=("Segoe UI", f), background=LIGHT_BG, foreground=PRIMARY, borderwidth=1, bordercolor=BORDER_COLOR)
        self.style.configure("TEntry", fieldbackground="#ffffff", foreground="#000000", insertcolor="#0f172a", insertwidth=2)
        self.style.configure("TSpinbox", fieldbackground="#ffffff", foreground="#000000", insertcolor="#0f172a", insertwidth=2)
        self.style.configure("TCombobox", fieldbackground="#ffffff", foreground="#000000", insertcolor="#0f172a", insertwidth=2)
        self.style.configure("Accent.TButton", font=("Segoe UI", f + 1, "bold"), background=ACCENT, foreground="#ffffff", borderwidth=0)
        self.style.configure("Export.TButton", font=("Segoe UI", f), background="#0f766e", foreground="#ffffff", borderwidth=0)
        self.style.map("Export.TButton", background=[("active", "#115e59")])
        self.style.configure("Treeview", font=("Segoe UI", f), background=BG_COLOR, foreground="#334155", rowheight=int(f * 2.4))
        self.style.configure("Treeview.Heading", font=("Segoe UI", f, "bold"), background="#f1f5f9", foreground=PRIMARY)

    # ----------------------------------------------------------------------
    # TAB 1: APERTURA I CANVI DE TORN
    # ----------------------------------------------------------------------
    def crear_tab_horaris(self):
        f_top = ttk.Frame(self.tab_horaris)
        f_top.pack(fill="x", pady=(15, 5), padx=20)

        ttk.Label(f_top, text="Horari d'Apertura General (setmana tipus, es repeteix cada setmana)",
                  font=("Segoe UI", 13, "bold")).pack(side="left")
        ic_copiar = icona_emoji("📋", 16); self.icones.append(ic_copiar)
        ttk.Button(f_top, text=" Copiar Dilluns a Tota la Setmana", image=ic_copiar, compound="left", command=self.copiar_dilluns_a_tots).pack(side="right", padx=5)

        f_franja = ttk.LabelFrame(self.tab_horaris, text=" Franja de Canvi de Torn ", padding=10)
        f_franja.pack(fill="x", padx=20, pady=10)

        ttk.Label(f_franja, text="Canvi de torn entre:").grid(row=0, column=0, sticky="w")
        ttk.Entry(f_franja, textvariable=self.inici_canvi_var, width=8).grid(row=0, column=1, padx=5)
        ttk.Label(f_franja, text="i").grid(row=0, column=2)
        ttk.Entry(f_franja, textvariable=self.fi_canvi_var, width=8).grid(row=0, column=3, padx=5)

        ttk.Label(f_franja, text="Durada mín. torn (h):").grid(row=1, column=0, sticky="w", pady=(10, 0))
        ttk.Entry(f_franja, textvariable=self.durada_min_torn_var, width=8).grid(row=1, column=1, padx=5, pady=(10, 0))
        ttk.Label(f_franja, text="Durada màx. torn (h):").grid(row=1, column=2, sticky="w", padx=(15, 5), pady=(10, 0))
        ttk.Entry(f_franja, textvariable=self.durada_max_torn_var, width=8).grid(row=1, column=3, padx=5, pady=(10, 0))

        ttk.Label(
            f_franja,
            text="ℹ️ Un fragment més curt que la durada mínima es fusiona en un únic torn de tot el dia.",
            font=("Segoe UI", 9, "italic"), foreground="#64748b"
        ).grid(row=2, column=0, columnspan=5, sticky="w", pady=(8, 0))

        frame_taula = ttk.Frame(self.tab_horaris)
        frame_taula.pack(fill="x", padx=20, pady=10)

        capçaleres = ["Dia", "Hora Obertura", "Hora Tancament", "Accions"]
        for col, text in enumerate(capçaleres):
            ttk.Label(frame_taula, text=text, font=("Segoe UI", 10, "bold")).grid(row=0, column=col, padx=15, pady=8, sticky="w")

        for i, dia in enumerate(self.dies):
            ttk.Label(frame_taula, text=dia, font=("Segoe UI", 10, "bold")).grid(row=i + 1, column=0, padx=15, pady=6, sticky="w")

            inici = tk.StringVar(value="08:30" if i < 5 else ("09:00" if i == 5 else ""))
            fi = tk.StringVar(value="20:30" if i < 5 else ("14:00" if i == 5 else ""))
            ttk.Entry(frame_taula, textvariable=inici, width=12).grid(row=i + 1, column=1, padx=15)
            ttk.Entry(frame_taula, textvariable=fi, width=12).grid(row=i + 1, column=2, padx=15)

            ttk.Button(frame_taula, text="Buidar", command=lambda d=dia: self.buidar_dia(d), width=8).grid(row=i + 1, column=3, padx=15)
            self.horaris_ui[dia] = {'inici': inici, 'fi': fi}

    def copiar_dilluns_a_tots(self):
        i_val = self.horaris_ui['Dilluns']['inici'].get()
        f_val = self.horaris_ui['Dilluns']['fi'].get()
        for dia in self.dies[1:5]:
            self.horaris_ui[dia]['inici'].set(i_val)
            self.horaris_ui[dia]['fi'].set(f_val)

    def buidar_dia(self, dia):
        self.horaris_ui[dia]['inici'].set("")
        self.horaris_ui[dia]['fi'].set("")

    # ----------------------------------------------------------------------
    # TAB 2: EQUIP, PATRÓ SETMANAL I ROTACIÓ
    # ----------------------------------------------------------------------
    def crear_tab_treb(self):
        f_top = ttk.LabelFrame(self.tab_treb, text=" Restriccions Generals de Personal ", padding=10)
        f_top.pack(fill="x", padx=20, pady=5)

        ttk.Label(f_top, text="Min/Màx Personal Torn (Dl-Dv):").grid(row=0, column=0, sticky="w")
        ttk.Entry(f_top, textvariable=self.min_personal_setmana_var, width=5).grid(row=0, column=1, padx=5)
        ttk.Label(f_top, text="-").grid(row=0, column=2)
        ttk.Entry(f_top, textvariable=self.max_personal_setmana_var, width=5).grid(row=0, column=3, padx=5)

        ttk.Label(f_top, text="Llicenciats mínims/Torn (Dl-Dv):").grid(row=0, column=4, sticky="w", padx=(20, 5))
        ttk.Entry(f_top, textvariable=self.min_llicenciats_var, width=5).grid(row=0, column=5)      

        ttk.Label(f_top, text="Min/Màx Personal Torn (Ds-Dg):").grid(row=1, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(f_top, textvariable=self.min_personal_cap_setmana_var, width=5).grid(row=1, column=1, padx=5, pady=(8, 0))
        ttk.Label(f_top, text="-").grid(row=1, column=2, pady=(8, 0))
        ttk.Entry(f_top, textvariable=self.max_personal_cap_setmana_var, width=5).grid(row=1, column=3, padx=5, pady=(8, 0))

        ttk.Label(f_top, text="Llicenciats mínims/Torn (Ds-Dg):").grid(row=1, column=4, sticky="w", padx=(20, 5), pady=(8, 0))
        ttk.Entry(f_top, textvariable=self.min_llicenciats_cap_setmana_var, width=5).grid(row=1, column=5, pady=(8, 0))

        ttk.Label(
            f_top,
            text="ℹ️ Sense patró fix de cap de setmana → grup de ROTACIÓ (repartit equitativament, màx. 1 cada X setmanes). "
                 "Amb patró fix → exempt. Còmput d'hores per cicle (h/setmana × X setmanes).",
            font=("Segoe UI", 9, "italic"), foreground="#64748b", wraplength=1200, justify="left"
        ).grid(row=2, column=0, columnspan=8, sticky="w", pady=(8, 0))

        f_treb = ttk.LabelFrame(self.tab_treb, text=" Alta, Edició i Llista de Personal ", padding=10)
        f_treb.pack(fill="both", expand=True, padx=20, pady=5)

        self.nom_var = tk.StringVar()
        self.hores_var = tk.DoubleVar(value=40.0)
        self.llicenciat_var = tk.BooleanVar()
        self.extres_var = tk.BooleanVar(value=False)
        self.actiu_var = tk.BooleanVar(value=True)
        self.garantir_cs_var = tk.BooleanVar(value=False)

        f_form = ttk.Frame(f_treb)
        f_form.pack(fill="x", pady=5)

        ttk.Label(f_form, text="Nom:").grid(row=0, column=0, padx=4)
        ttk.Entry(f_form, textvariable=self.nom_var, width=12).grid(row=0, column=1, padx=4)
        ttk.Label(f_form, text="Hores/Setmana:").grid(row=0, column=2, padx=4)
        ttk.Entry(f_form, textvariable=self.hores_var, width=5).grid(row=0, column=3, padx=4)

        ttk.Checkbutton(f_form, text="Llicenciat/da", variable=self.llicenciat_var).grid(row=0, column=4, padx=6)
        ttk.Checkbutton(f_form, text="Permetre Extres", variable=self.extres_var).grid(row=0, column=5, padx=6)
        ttk.Checkbutton(f_form, text="D'alta (Actiu/va)", variable=self.actiu_var).grid(row=0, column=6, padx=6)

        ic_garantir = icona_emoji("🔁", 14); self.icones.append(ic_garantir)
        ttk.Checkbutton(f_form, text=" Garantir ≥1 cap de setmana al cicle", image=ic_garantir, compound="left", variable=self.garantir_cs_var).grid(row=1, column=0, columnspan=4, sticky="w", padx=4, pady=(6, 0))

        ic_afegir = icona_emoji("➕", 16); self.icones.append(ic_afegir)
        ic_actualitzar = icona_emoji("💾", 16); self.icones.append(ic_actualitzar)
        ic_eliminar = icona_emoji("🗑️", 16); self.icones.append(ic_eliminar)
        ttk.Button(f_form, text=" Afegir", image=ic_afegir, compound="left", command=self.afegir_treballador).grid(row=0, column=7, padx=3)
        ttk.Button(f_form, text=" Actualitzar", image=ic_actualitzar, compound="left", command=self.modificar_treballador).grid(row=0, column=8, padx=3)
        ttk.Button(f_form, text=" Eliminar", image=ic_eliminar, compound="left", command=self.eliminar_treballador).grid(row=0, column=9, padx=3)

        self.tree_treb = ttk.Treeview(f_treb, columns=("Nom", "Hores", "Llicenciat", "Extres", "Estat", "Garantit"), show='headings', height=6)
        self.tree_treb.heading("Nom", text="Nom")
        self.tree_treb.heading("Hores", text="Hores/Setmana")
        self.tree_treb.heading("Llicenciat", text="Llicenciat/da")
        self.tree_treb.heading("Extres", text="Extres")
        self.tree_treb.heading("Estat", text="Estat")
        self.tree_treb.heading("Garantit", text="🔁 Garantit CS")

        self.tree_treb.column("Nom", width=150)
        self.tree_treb.column("Hores", width=100, anchor="center")
        self.tree_treb.column("Llicenciat", width=100, anchor="center")
        self.tree_treb.column("Extres", width=90, anchor="center")
        self.tree_treb.column("Estat", width=110, anchor="center")
        self.tree_treb.column("Garantit", width=100, anchor="center")

        self.tree_treb.pack(fill="both", expand=True, pady=5)
        self.tree_treb.bind("<<TreeviewSelect>>", self.carregar_treballador_seleccionat)
        self._configurar_zebra(self.tree_treb)

        f_patro = ttk.LabelFrame(self.tab_treb, text=" 📅 Horari Setmanal Tipus ", padding=10)
        f_patro.pack(fill="x", padx=20, pady=(5, 10))

        f_header_patro = ttk.Frame(f_patro)
        f_header_patro.pack(fill="x", pady=(0, 6))

        self.lbl_treb_seleccionat = ttk.Label(f_header_patro, text="👉 Selecciona un treballador per definir la seva setmana",
                                               font=("Segoe UI", 10, "bold"), foreground="#2563eb")
        self.lbl_treb_seleccionat.pack(side="left")

        f_actions = ttk.Frame(f_header_patro)
        f_actions.pack(side="right")
        ic_mati = icona_emoji("🌅", 16); self.icones.append(ic_mati)
        ic_tarde = icona_emoji("🌆", 16); self.icones.append(ic_tarde)
        ic_reset = icona_emoji("🔄", 16); self.icones.append(ic_reset)
        ttk.Button(f_actions, text=" Tot Matí (L-V)", image=ic_mati, compound="left", command=self.patro_tot_mati).pack(side="left", padx=2)
        ttk.Button(f_actions, text=" Tot Tarde (L-V)", image=ic_tarde, compound="left", command=self.patro_tot_tarde).pack(side="left", padx=2)
        ttk.Button(f_actions, text=" Tot Indiferent", image=ic_reset, compound="left", command=self.patro_reset).pack(side="left", padx=2)
        f_grid = ttk.Frame(f_patro)
        f_grid.pack(fill="x")

        for col, dia in enumerate(self.dies):
            f_col = ttk.Frame(f_grid, padding=4, borderwidth=1, relief="solid")
            f_col.grid(row=0, column=col, padx=3, pady=2, sticky="nsew")
            f_grid.columnconfigure(col, weight=1)

            ttk.Label(f_col, text=dia, font=("Segoe UI", 10, "bold"), anchor="center").pack(fill="x", pady=(0, 5))

            cb = ttk.Combobox(f_col, values=OPCIONS_PATRO, state="readonly", width=15)
            cb.set("-")
            cb.pack(fill="x", pady=2)
            cb.bind("<<ComboboxSelected>>", self.guardar_patro_ui)
            self.combos_patro[dia] = cb

        ttk.Label(
            f_patro,
            text="💡 Ds/Dg en '-' → grup de rotació. Fixats explícitament → exempt.",
            font=("Segoe UI", 9, "italic"), foreground="#64748b"
        ).pack(fill="x", pady=(8, 0))

    def patro_tot_mati(self):
        nom = self.nom_var.get()
        if not nom or nom not in self.treballadors:
            return
        for dia in self.dies[:5]:
            self.combos_patro[dia].set("Preferentment Matí")
        self.guardar_patro_ui()

    def patro_tot_tarde(self):
        nom = self.nom_var.get()
        if not nom or nom not in self.treballadors:
            return
        for dia in self.dies[:5]:
            self.combos_patro[dia].set("Preferentment Tarda")
        self.guardar_patro_ui()

    def patro_reset(self):
        nom = self.nom_var.get()
        if not nom or nom not in self.treballadors:
            return
        for dia in self.dies:
            self.combos_patro[dia].set("-")
        self.guardar_patro_ui()

    def carregar_treballador_seleccionat(self, event=None):
        sel = self.tree_treb.selection()
        if not sel:
            return
        item = self.tree_treb.item(sel[0])
        nom = item['values'][0]
        info = self.treballadors[nom]
        self.nom_var.set(nom)
        self.hores_var.set(info['hores'])
        self.llicenciat_var.set(info['llicenciat'])
        self.extres_var.set(info.get('extres', False))
        self.actiu_var.set(info.get('actiu', True))
        self.garantir_cs_var.set(info.get('garantir_cap_setmana', False))
        self.lbl_treb_seleccionat.config(text=f"👉 Editant patró setmanal de: {nom}")

        patro = info.get('patro_semanal', {})
        for dia in self.dies:
            val = patro.get(dia, "-")
            self.combos_patro[dia].set(val)

    def guardar_patro_ui(self, event=None):
        nom = self.nom_var.get()
        if not nom or nom not in self.treballadors:
            return

        if 'patro_semanal' not in self.treballadors[nom]:
            self.treballadors[nom]['patro_semanal'] = {}

        for dia in self.dies:
            val = self.combos_patro[dia].get()
            self.treballadors[nom]['patro_semanal'][dia] = val

    def afegir_treballador(self):
        nom = self.nom_var.get().strip()
        if not nom:
            messagebox.showwarning("Atenció", "Introdueix un nom de treballador vàlid.")
            return
        if nom in self.treballadors:
            messagebox.showwarning("Atenció", "Aquest treballador ja existeix. Utilitza 'Actualitzar'.")
            return

        patro_actual = {dia: self.combos_patro[dia].get() for dia in self.dies}
        self.treballadors[nom] = {
            'hores': self.hores_var.get(),
            'llicenciat': self.llicenciat_var.get(),
            'extres': self.extres_var.get(),
            'actiu': self.actiu_var.get(),
            'garantir_cap_setmana': self.garantir_cs_var.get(),
            'patro_semanal': patro_actual
        }

        self.actualitzar_llista_treballadors()
        self.buidar_formulari_treballador()

    def modificar_treballador(self):
        nom = self.nom_var.get().strip()
        if not nom or nom not in self.treballadors:
            messagebox.showwarning("Atenció", "Selecciona un treballador existent per modificar.")
            return

        patro_actual = {dia: self.combos_patro[dia].get() for dia in self.dies}
        self.treballadors[nom] = {
            'hores': self.hores_var.get(),
            'llicenciat': self.llicenciat_var.get(),
            'extres': self.extres_var.get(),
            'actiu': self.actiu_var.get(),
            'garantir_cap_setmana': self.garantir_cs_var.get(),
            'patro_semanal': patro_actual
        }

        self.actualitzar_llista_treballadors()

    def eliminar_treballador(self):
        nom = self.nom_var.get().strip()
        if nom in self.treballadors:
            del self.treballadors[nom]
        self.actualitzar_llista_treballadors()
        self.buidar_formulari_treballador()

    def buidar_formulari_treballador(self):
        self.nom_var.set("")
        self.hores_var.set(40.0)
        self.llicenciat_var.set(False)
        self.extres_var.set(False)
        self.actiu_var.set(True)
        self.garantir_cs_var.set(False)
        self.lbl_treb_seleccionat.config(text="👉 Selecciona un treballador per definir la seva setmana")
        for dia in self.dies:
            self.combos_patro[dia].set("-")

    def actualitzar_llista_treballadors(self):
        for row in self.tree_treb.get_children():
            self.tree_treb.delete(row)
        for nom, info in self.treballadors.items():
            estat_txt = "🟢 Alta" if info.get('actiu', True) else "🔴 Baixa"
            self.tree_treb.insert("", "end", values=(
                nom,
                info['hores'],
                "Sí" if info['llicenciat'] else "No",
                "Sí" if info.get('extres', False) else "No",
                estat_txt,
                "Sí" if info.get('garantir_cap_setmana', False) else "No"
            ))
        self._aplicar_zebra(self.tree_treb)

    # ----------------------------------------------------------------------
    # TAB 3: HORARI CÍCLIC OPTIMITZAT
    # ----------------------------------------------------------------------
    def crear_tab_resultats(self):
        f_top = ttk.Frame(self.tab_resultats)
        f_top.pack(fill="x", padx=20, pady=10)

        ic_generar = icona_emoji("🚀", 20); self.icones.append(ic_generar)
        ic_pdf = icona_emoji("📄", 18); self.icones.append(ic_pdf)
        ic_excel = icona_emoji("📊", 18); self.icones.append(ic_excel)

        btn = ttk.Button(f_top, text=" GENERAR HORARI CÍCLIC", image=ic_generar, compound="left",
            style="Accent.TButton", command=self.optimitzar_i_mostrar)
        btn.pack(side="left", padx=5, ipady=5)

        ttk.Button(f_top, text=" Exportar PDF", image=ic_pdf, compound="left",
            style="Export.TButton", command=self.exportar_pdf).pack(side="left", padx=5, ipady=5)
        ttk.Button(f_top, text=" Exportar Excel", image=ic_excel, compound="left",
            style="Export.TButton", command=self.exportar_excel).pack(side="left", padx=5, ipady=5)

        ttk.Label(f_top, text="Setmana:", font=("Segoe UI", 10, "bold")).pack(side="left", padx=(15, 5))
        self.cb_setmana = ttk.Combobox(f_top, state="readonly", width=30)
        self.cb_setmana.pack(side="left", padx=5)
        self.cb_setmana.bind("<<ComboboxSelected>>", self.actualitzar_vista_setmanal)

        ttk.Label(f_top, text="Vista:", font=("Segoe UI", 10, "bold")).pack(side="left", padx=(20, 5))
        self.cb_vista = ttk.Combobox(f_top, values=["Per Treballador", "Per Torn (Matí / Tarda)"], state="readonly", width=22)
        self.cb_vista.current(0)
        self.cb_vista.pack(side="left", padx=5)
        self.cb_vista.bind("<<ComboboxSelected>>", self.actualitzar_vista_setmanal)

        self.f_info_franges = ttk.LabelFrame(self.tab_resultats, text=" Resum ", padding=(10, 6))
        self.f_info_franges.pack(fill="x", padx=20, pady=(0, 8))

        self.lbl_franges_text = ttk.Label(
            self.f_info_franges, text="Fes clic a 'Generar Horari Cíclic'.",
            font=("Segoe UI", 10, "bold"), foreground="#1e293b"
        )
        self.lbl_franges_text.pack(anchor="w")

        self.lbl_avisos_cap_setmana = ttk.Label(
            self.f_info_franges, text="", font=("Segoe UI", 9, "italic"), foreground="#b45309",
            wraplength=1300, justify="left"
        )
        self.lbl_avisos_cap_setmana.pack(anchor="w", pady=(4, 0))

        self.lbl_avisos_llicenciats = ttk.Label(
            self.f_info_franges, text="", font=("Segoe UI", 9, "bold"), foreground="#dc2626",
            wraplength=1300, justify="left"
        )
        self.lbl_avisos_llicenciats.pack(anchor="w", pady=(4, 0))

        self.canvas_horari_wrap = ttk.Frame(self.tab_resultats)
        self.canvas_horari_wrap.pack(fill="both", expand=True, padx=20, pady=(0, 5))

        self.canvas_horari = tk.Canvas(self.canvas_horari_wrap, bg="#ffffff", highlightthickness=0)
        self.scroll_v_horari = ttk.Scrollbar(self.canvas_horari_wrap, orient="vertical", command=self.canvas_horari.yview)
        self.scroll_h_horari = ttk.Scrollbar(self.canvas_horari_wrap, orient="horizontal", command=self.canvas_horari.xview)
        self.canvas_horari.configure(yscrollcommand=self.scroll_v_horari.set, xscrollcommand=self.scroll_h_horari.set)

        self.canvas_horari.grid(row=0, column=0, sticky="nsew")
        self.scroll_v_horari.grid(row=0, column=1, sticky="ns")
        self.scroll_h_horari.grid(row=1, column=0, sticky="ew")
        self.canvas_horari_wrap.rowconfigure(0, weight=1)
        self.canvas_horari_wrap.columnconfigure(0, weight=1)

        self.frame_grid_horari = tk.Frame(self.canvas_horari, bg="#ffffff")
        self._id_frame_grid_horari = self.canvas_horari.create_window((0, 0), window=self.frame_grid_horari, anchor="nw")
        self.frame_grid_horari.bind(
            "<Configure>",
            lambda e: self.canvas_horari.configure(scrollregion=self.canvas_horari.bbox("all"))
        )

        self.f_resum_frame = ttk.LabelFrame(self.tab_resultats, text=" Balanç d'Hores per Cicle ", padding=10)
        self.f_resum_frame.pack(fill="x", padx=20, pady=(5, 10))

        self.tree_resum_cicle = ttk.Treeview(
            self.f_resum_frame,
            columns=("Nom", "Contracte", "HoresSetmanaSel", "ObjSetmana", "BalancSetmana", "ObjCicle", "Estat"),
            show='headings',
            height=3
        )

        ic_h_contracte = icona_emoji("📋", 14); self.icones.append(ic_h_contracte)
        ic_h_fetes = icona_emoji("📅", 14); self.icones.append(ic_h_fetes)
        ic_h_obj1 = icona_emoji("🎯", 14); self.icones.append(ic_h_obj1)
        ic_h_balanc = icona_emoji("⚖️", 14); self.icones.append(ic_h_balanc)
        ic_h_obj2 = icona_emoji("🎯", 14); self.icones.append(ic_h_obj2)
        ic_h_estat = icona_emoji("📌", 14); self.icones.append(ic_h_estat)

        self.tree_resum_cicle.heading("Nom", text="Treballador/a")
        self.tree_resum_cicle.heading("Contracte", text=" Contracte (h/s)", image=ic_h_contracte)
        self.tree_resum_cicle.heading("HoresSetmanaSel", text=" Fetes (setmana)", image=ic_h_fetes)
        self.tree_resum_cicle.heading("ObjSetmana", text=" Obj. Setmana", image=ic_h_obj1)
        self.tree_resum_cicle.heading("BalancSetmana", text=" Balanç", image=ic_h_balanc)
        self.tree_resum_cicle.heading("ObjCicle", text=" Obj. Cicle", image=ic_h_obj2)
        self.tree_resum_cicle.heading("Estat", text=" Estat del Cicle", image=ic_h_estat)

        self.tree_resum_cicle.column("Nom", width=140, anchor="w")
        self.tree_resum_cicle.column("Contracte", width=110, anchor="center")
        self.tree_resum_cicle.column("HoresSetmanaSel", width=130, anchor="center")
        self.tree_resum_cicle.column("ObjSetmana", width=140, anchor="center")
        self.tree_resum_cicle.column("BalancSetmana", width=130, anchor="center")
        self.tree_resum_cicle.column("ObjCicle", width=110, anchor="center")
        self.tree_resum_cicle.column("Estat", width=220, anchor="center")

        self.tree_resum_cicle.pack(fill="x", expand=True)
        self._configurar_zebra(self.tree_resum_cicle)

    # ----------------------------------------------------------------------
    # LÒGICA AUXILIAR
    # ----------------------------------------------------------------------
    def es_dia_cap_setmana(self, dia_nom):
        return dia_nom in DIES_CAP_SETMANA

    def te_patro_fix_cap_setmana(self, info):
        patro = info.get('patro_semanal', {})
        sat_val = patro.get('Dissabte', '-')
        sun_val = patro.get('Diumenge', '-')
        return (sat_val != '-') or (sun_val != '-')

    def es_treballador_nomes_cap_setmana(self, info):
        patro = info.get('patro_semanal', {})
        dies_feiners = self.dies[:5]
        tots_no_ve_feiners = all(patro.get(d, "-") == "No ve" for d in dies_feiners)
        if not tots_no_ve_feiners:
            return False
        disponible_cap_setmana = any(
            patro.get(d, "-") not in ("-", "No ve") for d in DIES_CAP_SETMANA
        )
        return disponible_cap_setmana

    def min_max_personal_per_dia(self, dia_nom):
        if self.es_dia_cap_setmana(dia_nom):
            return self.min_personal_cap_setmana_var.get(), self.max_personal_cap_setmana_var.get()
        return self.min_personal_setmana_var.get(), self.max_personal_setmana_var.get()

    def min_llicenciats_per_dia(self, dia_nom):
        if self.es_dia_cap_setmana(dia_nom):
            return self.min_llicenciats_cap_setmana_var.get()
        return self.min_llicenciats_var.get()

    def calcular_torns_del_dia(self, h_inici, h_fi, inici_canvi_f, fi_canvi_f):
        durada_min = max(0.0, self.durada_min_torn_var.get())
        durada_max = self.durada_max_torn_var.get()

        pausa_mid = (inici_canvi_f + fi_canvi_f) / 2.0
        durada_m = pausa_mid - h_inici
        durada_t = h_fi - pausa_mid

        excedeix_max = False

        if durada_m > 0 and durada_t > 0:
            if durada_m < durada_min or durada_t < durada_min:
                durada_total = h_fi - h_inici
                if durada_total > durada_max:
                    excedeix_max = True
                torns = [("Matí", durada_total, f"{to_hhmm(h_inici)}-{to_hhmm(h_fi)}")]
                return torns, to_hhmm(h_fi), excedeix_max
            if durada_m > durada_max or durada_t > durada_max:
                excedeix_max = True
            torns = [
                ("Matí", durada_m, f"{to_hhmm(h_inici)}-{to_hhmm(pausa_mid)}"),
                ("Tarda", durada_t, f"{to_hhmm(pausa_mid)}-{to_hhmm(h_fi)}")
            ]
            return torns, to_hhmm(pausa_mid), excedeix_max
        elif durada_m > 0:
            if durada_m > durada_max:
                excedeix_max = True
            return [("Matí", durada_m, f"{to_hhmm(h_inici)}-{to_hhmm(h_fi)}")], to_hhmm(h_fi), excedeix_max
        elif durada_t > 0:
            if durada_t > durada_max:
                excedeix_max = True
            return [("Tarda", durada_t, f"{to_hhmm(h_inici)}-{to_hhmm(h_fi)}")], to_hhmm(h_inici), excedeix_max
        else:
            return [], to_hhmm(pausa_mid), False

    def calcular_torns_per_dia_nom(self):
        torns_per_dia_nom = {}
        h_canvi_detectada = f"{self.inici_canvi_var.get()}-{self.fi_canvi_var.get()}"
        avisos_durada = []

        try:
            inici_canvi_f = float(self.inici_canvi_var.get().split(':')[0]) + float(self.inici_canvi_var.get().split(':')[1]) / 60.0
            fi_canvi_f = float(self.fi_canvi_var.get().split(':')[0]) + float(self.fi_canvi_var.get().split(':')[1]) / 60.0
        except Exception:
            inici_canvi_f, fi_canvi_f = 13.5, 14.5

        for dia_nom in self.dies:
            inici_str = self.horaris_ui[dia_nom]['inici'].get().strip()
            fi_str = self.horaris_ui[dia_nom]['fi'].get().strip()

            if not inici_str or not fi_str:
                torns_per_dia_nom[dia_nom] = []
                continue

            try:
                h_inici = float(inici_str.split(':')[0]) + float(inici_str.split(':')[1]) / 60.0
                h_fi = float(fi_str.split(':')[0]) + float(fi_str.split(':')[1]) / 60.0
            except Exception:
                torns_per_dia_nom[dia_nom] = []
                continue

            torns, hora_canvi_dia, excedeix_max = self.calcular_torns_del_dia(h_inici, h_fi, inici_canvi_f, fi_canvi_f)
            torns_per_dia_nom[dia_nom] = torns
            if len(torns) > 1:
                h_canvi_detectada = hora_canvi_dia
            if excedeix_max:
                avisos_durada.append(f"• {dia_nom}: un torn supera la durada màxima configurada ({self.durada_max_torn_var.get():.1f} h).")

        return torns_per_dia_nom, h_canvi_detectada, avisos_durada

    def analitzar_infactibilitat(self, torns_per_dia_nom):
        errors = []
        treballadors_actius = {t: info for t, info in self.treballadors.items() if info.get('actiu', True)}

        for dia_nom in self.dies:
            if not torns_per_dia_nom.get(dia_nom):
                continue
            min_p, _ = self.min_max_personal_per_dia(dia_nom)

            for nt, _, _ in torns_per_dia_nom[dia_nom]:
                disponibles = []
                for t, info in treballadors_actius.items():
                    patro = info.get('patro_semanal', {})
                    val = patro.get(dia_nom, "-")

                    pot = True
                    if val in ["No ve", "Festa"]:
                        pot = False
                    elif val == "Només Matí" and nt == "Tarda":
                        pot = False
                    elif val == "Només Tarda" and nt == "Matí":
                        pot = False

                    if pot:
                        disponibles.append(t)

                if len(disponibles) < min_p:
                    errors.append(f"• {dia_nom} - Torn {nt}: Només hi ha {len(disponibles)} disponible(s), cal un mínim de {min_p}.")

        return errors

    def analitzar_capacitat_individual_cicle(self, torns_per_dia_nom, num_setmanes, target_cicle_dict):
        errors = []
        treballadors_actius = {t: info for t, info in self.treballadors.items() if info.get('actiu', True)}

        for t, info in treballadors_actius.items():
            objectiu = target_cicle_dict.get(t, 0.0)
            if objectiu <= 0.01:
                continue

            patro = info.get('patro_semanal', {})
            max_hores_una_setmana = 0.0
            for dia_nom in self.dies:
                val = patro.get(dia_nom, "-")
                if val in ("No ve", "Festa"):
                    continue
                for nt, durada, _ in torns_per_dia_nom.get(dia_nom, []):
                    if val == "Només Matí" and nt == "Tarda":
                        continue
                    if val == "Només Tarda" and nt == "Matí":
                        continue
                    max_hores_una_setmana += durada

            if not self.te_patro_fix_cap_setmana(info):
                hores_cap_setmana_disponibles = 0.0
                for dia_nom in DIES_CAP_SETMANA:
                    val = patro.get(dia_nom, "-")
                    if val in ("No ve", "Festa"):
                        continue
                    for nt, durada, _ in torns_per_dia_nom.get(dia_nom, []):
                        hores_cap_setmana_disponibles += durada
                max_hores_setmana_normal = max_hores_una_setmana - hores_cap_setmana_disponibles
                max_hores_disponibles = max_hores_setmana_normal * num_setmanes + hores_cap_setmana_disponibles
            else:
                max_hores_disponibles = max_hores_una_setmana * num_setmanes

            if max_hores_disponibles < objectiu - 0.05:
                falten = objectiu - max_hores_disponibles
                errors.append(
                    f"• {t}: hauria de fer {objectiu:.1f}h en aquest cicle de {num_setmanes} setmana(es), però amb el "
                    f"seu patró setmanal actual com a màxim hi pot arribar a {max_hores_disponibles:.1f}h. "
                    f"Li falten {falten:.1f}h. → Amplia la seva disponibilitat, redueix les hores de contracte, o "
                    f"allarga l'horari d'obertura d'algun dia."
                )

        return errors

    def analitzar_capacitat_cap_setmana(self, torns_per_dia_nom, num_setmanes):
        errors = []
        treballadors_actius = {t: info for t, info in self.treballadors.items() if info.get('actiu', True)}
        min_p_cs = self.min_personal_cap_setmana_var.get()

        hi_ha_cs_obert = any(torns_per_dia_nom.get(d) for d in DIES_CAP_SETMANA)
        if not hi_ha_cs_obert:
            return errors

        no_exempts = {t: info for t, info in treballadors_actius.items() if not self.te_patro_fix_cap_setmana(info)}
        exempts = {t: info for t, info in treballadors_actius.items() if self.te_patro_fix_cap_setmana(info)}

        te_algu_garantit_cada_setmana = False
        for t, info in exempts.items():
            patro = info.get('patro_semanal', {})
            for dia_nom in DIES_CAP_SETMANA:
                val = patro.get(dia_nom, "-")
                if val in ("Tot el dia", "Només aquest dia"):
                    te_algu_garantit_cada_setmana = True

        necessitat_personal = num_setmanes * min_p_cs
        capacitat_rotacio = len(no_exempts) * 1
        garantia_personal = num_setmanes if te_algu_garantit_cada_setmana else 0

        if garantia_personal + capacitat_rotacio < necessitat_personal:
            errors.append(
                f"• PERSONAL AL CAP DE SETMANA: calen {necessitat_personal} assignacions en tot el cicle "
                f"({num_setmanes} setmana(es) × mínim {min_p_cs} persones), però amb la regla '1 cap de setmana cada "
                f"{num_setmanes} setmanes' només hi ha capacitat per {garantia_personal + capacitat_rotacio}. "
                f"→ Escurça el cicle, dona d'alta més personal disponible el cap de setmana, o defineix patrons fixos."
            )

        return errors

    def analitzar_garantits_impossibles(self, num_setmanes):
        """
        Comprova que els treballadors marcats amb 'garantir ≥1 cap de setmana'
        no superin la capacitat màxima de caps de setmana disponibles al cicle.
        """
        errors = []
        treballadors_actius = {t: info for t, info in self.treballadors.items() if info.get('actiu', True)}
        garantits = [t for t, info in treballadors_actius.items()
                     if info.get('garantir_cap_setmana', False) and not self.te_patro_fix_cap_setmana(info)]
        max_cs = self.max_personal_cap_setmana_var.get()
        if len(garantits) > max_cs * num_setmanes:
            errors.append(
                f"• GARANTITS IMPOSSIBLES: hi ha {len(garantits)} treballador(s) marcats per garantir un cap de "
                f"setmana, però la capacitat total del cicle és de {max_cs * num_setmanes} places "
                f"(màx. personal cap de setmana × núm. setmanes). → Redueix garantits, puja el màxim, o allarga el cicle."
            )
        return errors

    # ----------------------------------------------------------------------
    # GENERACIÓ / OPTIMITZACIÓ
    # ----------------------------------------------------------------------
    def optimitzar_i_mostrar(self):
        try:
            num_setmanes = max(1, self.num_setmanes_cicle_var.get())

            res_quadrant, res_resum, torns_per_dia_nom, target_setmana_dict = self.resolver_model_cicle(num_setmanes)

            if not res_quadrant:
                messagebox.showerror("Error de Conjunt de Restriccions", res_resum)
                return

            self.res_quadrant_cicle = res_quadrant
            self.res_resum_cicle = res_resum
            self.torns_per_dia_nom = torns_per_dia_nom
            self.num_setmanes_generat = num_setmanes
            self.target_setmana_dict = target_setmana_dict

            treballadors_actius = {t: info for t, info in self.treballadors.items() if info.get('actiu', True)}

            noms_nomes_cap_setmana = [t for t, info in treballadors_actius.items() if self.es_treballador_nomes_cap_setmana(info)]
            self.avisos_cap_setmana = noms_nomes_cap_setmana
            if noms_nomes_cap_setmana:
                self.lbl_avisos_cap_setmana.config(
                    text="⚠️ Només CS (exempts de rotació): " + ", ".join(noms_nomes_cap_setmana)
                )
            else:
                self.lbl_avisos_cap_setmana.config(text="")

            if self.dies_sense_llicenciat:
                dies_text = ", ".join(f"Setmana {w+1} - {d}" for (w, d) in sorted(self.dies_sense_llicenciat))
                self.lbl_avisos_llicenciats.config(
                    text=f"🚨 Sense mínim de llicenciats cobert: {dies_text}."
                )
            else:
                self.lbl_avisos_llicenciats.config(text="")

            opcions_setmanes = [f"Setmana {w + 1} de {num_setmanes} (cicle)" for w in range(num_setmanes)]
            self.cb_setmana['values'] = opcions_setmanes
            if opcions_setmanes:
                self.cb_setmana.current(0)
                self.actualitzar_vista_setmanal()

        except Exception as e:
            messagebox.showerror("Error de Generació", f"S'ha produït un error inesperat:\n{str(e)}")

    def actualitzar_vista_setmanal(self, event=None):
        if not self.res_quadrant_cicle:
            return

        idx_w = self.cb_setmana.current()
        if idx_w < 0:
            idx_w = 0

        mode_vista = self.cb_vista.get()

        inici_c = self.inici_canvi_var.get()
        fi_c = self.fi_canvi_var.get()
        hora_canvi = self.torns_per_dia_nom.get('hora_canvi', f"{inici_c}-{fi_c}")

        self.lbl_franges_text.config(
            text=f"Canvi de torn: {hora_canvi} h  |  Cicle de {self.num_setmanes_generat} setmana(es) (indefinit)"
        )

        for w in self.frame_grid_horari.winfo_children():
            w.destroy()

        hores_setmana_per_persona = {}
        fg = self.frame_grid_horari

        if mode_vista.startswith("Per Torn"):
            self._cella_horari(fg, "Torn / Horari", 0, 0, tipus="header", bold=True, ample=16)
            for c, dia_nom in enumerate(self.dies, start=1):
                self._cella_horari(fg, dia_nom, 0, c, tipus="header", bold=True, ample=22)

            for r, (etiqueta, nt_fix) in enumerate([("Matí", "Matí"), ("Tarda", "Tarda")], start=1):
                self._cella_horari(fg, etiqueta, r, 0, tipus=nt_fix, bold=True, ample=16)
                for c, dia_nom in enumerate(self.dies, start=1):
                    torns_dia = self.torns_per_dia_nom.get(dia_nom, [])
                    nts_disponibles = [nt for nt, _, _ in torns_dia]
                    dur = next((dur for nt, dur, _ in torns_dia if nt == nt_fix), 0)

                    if not torns_dia or nt_fix not in nts_disponibles:
                        self._cella_horari(fg, "-", r, c, tipus="-", ample=22)
                        continue

                    pers = []
                    for t, setmanes_dict in self.res_quadrant_cicle.items():
                        val = setmanes_dict.get(idx_w, {}).get(dia_nom, "Festa")
                        if (nt_fix == "Matí" and val in ["Matí", "Matí + Tarda"]) or \
                           (nt_fix == "Tarda" and val in ["Tarda", "Matí + Tarda"]): pers.append(t)
                        hores_setmana_per_persona[t] = hores_setmana_per_persona.get(t, 0.0) + dur

                    avis = (idx_w, dia_nom) in self.dies_sense_llicenciat
                    text_cel = ", ".join(pers) if pers else "—"
                    if avis:
                        text_cel += " (SENSE LLIC.)"
                    self._cella_horari(fg, text_cel, r, c, tipus=nt_fix, avis=avis, ample=22)

        else:
            capçalera = ["Treballador/a"] + self.dies + ["Total Setmana"]
            for c, txt in enumerate(capçalera):
                self._cella_horari(fg, txt, 0, c, tipus="header", bold=True, ample=16 if c == 0 else 14)

            for r, (t, setmanes_dict) in enumerate(self.res_quadrant_cicle.items(), start=1):
                self._cella_horari(fg, t, r, 0, tipus="-", bold=True, ample=16)
                h_setmana = 0.0
                dies_dict = setmanes_dict.get(idx_w, {})
                for c, dia_nom in enumerate(self.dies, start=1):
                    val = dies_dict.get(dia_nom, "Festa")
                    torns_dia = self.torns_per_dia_nom.get(dia_nom, [])

                    for nt, durada, _ in torns_dia:
                        if (val == "Matí" and nt == "Matí") or (val == "Tarda" and nt == "Tarda") or (val == "Matí + Tarda"):
                            h_setmana += durada

                    avis = (idx_w, dia_nom) in self.dies_sense_llicenciat
                    text_cel = val if val != "Matí + Tarda" else "M + T"
                    if avis:
                        text_cel += " (!)"
                    self._cella_horari(fg, text_cel, r, c, tipus=val, avis=avis, ample=14)

                hores_setmana_per_persona[t] = h_setmana
                self._cella_horari(fg, f"{h_setmana:.1f} h", r, len(self.dies) + 1, tipus="-", bold=True, ample=14)

        for row in self.tree_resum_cicle.get_children():
            self.tree_resum_cicle.delete(row)

        for t, dades in self.res_resum_cicle.items():
            h_e = dades['extres']
            h_d = dades.get('deficit', 0.0)
            h_contracte = self.treballadors[t]['hores']
            h_sel = hores_setmana_per_persona.get(t, 0.0)

            target_setmana = self.target_setmana_dict.get(t, h_contracte)
            balanc_setmana = h_sel - target_setmana
            objectiu_cicle = h_contracte * self.num_setmanes_generat

            balanc_str = f"+{balanc_setmana:.1f} h" if balanc_setmana > 0.05 else (f"{balanc_setmana:.1f} h" if balanc_setmana < -0.05 else "0.0 h")

            net_cicle = h_e - h_d
            if net_cicle > 0.05:
                estat_str = f"+{net_cicle:.1f}h de més"
            elif net_cicle < -0.05:
                estat_str = f"-{abs(net_cicle):.1f}h per sota"
            else:
                estat_str = "Hores mínimes complertes"

            if t in self.avisos_cap_setmana:
                estat_str += " · Només CS"

            self.tree_resum_cicle.insert("", "end", values=(
                t, f"{h_contracte:.1f} h/s", f"{h_sel:.1f} h", f"{target_setmana:.1f} h",
                balanc_str, f"{objectiu_cicle:.1f} h", estat_str
            ))

        self._aplicar_zebra(self.tree_resum_cicle)

    # ----------------------------------------------------------------------
    # MODEL D'OPTIMITZACIÓ (PULP) — CÍCLIC, SENSE MES NI FESTIUS
    # ----------------------------------------------------------------------
    def resolver_model_cicle(self, num_setmanes):
        treballadors_actius = {t: info for t, info in self.treballadors.items() if info.get('actiu', True)}
        if not treballadors_actius:
            return None, "❌ No hi ha cap treballador donat d'alta a la plantilla.", None, None

        torns_per_dia_nom, hora_canvi, avisos_durada = self.calcular_torns_per_dia_nom()
        torns_per_dia_nom['hora_canvi'] = hora_canvi

        target_setmana_dict = {t: info['hores'] for t, info in treballadors_actius.items()}
        target_cicle_dict = {t: info['hores'] * num_setmanes for t, info in treballadors_actius.items()}

        errors_diagnosticar = []
        if avisos_durada:
            errors_diagnosticar.append("DURADA DE TORN:")
            errors_diagnosticar.extend(avisos_durada)

        errors_capacitat_individual = self.analitzar_capacitat_individual_cicle(torns_per_dia_nom, num_setmanes, target_cicle_dict)
        if errors_capacitat_individual:
            errors_diagnosticar.append("HORES DE CONTRACTE INASSOLIBLES:")
            errors_diagnosticar.extend(errors_capacitat_individual)

        errors_diagnosticar.extend(self.analitzar_infactibilitat(torns_per_dia_nom))
        errors_diagnosticar.extend(self.analitzar_capacitat_cap_setmana(torns_per_dia_nom, num_setmanes))
        errors_diagnosticar.extend(self.analitzar_garantits_impossibles(num_setmanes))

        if errors_diagnosticar:
            msg = "❌ IMPOSSIBLE GENERAR L'HORARI\n\n" + "\n".join(errors_diagnosticar)
            return None, msg, None, None

        prob = pulp.LpProblem("Optimitzacio_Farmacia_Cicle", pulp.LpMinimize)

        x = {}
        for t in treballadors_actius:
            for w in range(num_setmanes):
                for dia_nom in self.dies:
                    for nt, _, _ in torns_per_dia_nom.get(dia_nom, []):
                        x[(t, w, dia_nom, nt)] = pulp.LpVariable(f"x_{t}_{w}_{dia_nom}_{nt}", cat='Binary')

        extra_h, deficit_h = {}, {}
        for t in treballadors_actius:
            extra_h[t] = pulp.LpVariable(f"extra_{t}", lowBound=0)
            deficit_h[t] = pulp.LpVariable(f"def_{t}", lowBound=0)
            prob += deficit_h[t] == 0

        funcio_obj = pulp.lpSum([1000 * extra_h[t] for t in treballadors_actius])

        for key, var in x.items():
            funcio_obj += random.uniform(SOROLL_ALEATORI_MIN, SOROLL_ALEATORI_MAX) * var

        for t, info in treballadors_actius.items():
            patro = info.get('patro_semanal', {})
            for w in range(num_setmanes):
                for dia_nom in self.dies:
                    val_especific = patro.get(dia_nom, "-")
                    if val_especific == "Preferentment Matí":
                        if (t, w, dia_nom, "Matí") in x:
                            funcio_obj += -50 * x[(t, w, dia_nom, "Matí")]
                        if (t, w, dia_nom, "Tarda") in x:
                            funcio_obj += 10 * x[(t, w, dia_nom, "Tarda")]
                    elif val_especific == "Preferentment Tarda":
                        if (t, w, dia_nom, "Tarda") in x:
                            funcio_obj += -50 * x[(t, w, dia_nom, "Tarda")]
                        if (t, w, dia_nom, "Matí") in x:
                            funcio_obj += 10 * x[(t, w, dia_nom, "Matí")]

        manca_llic = {}
        for w in range(num_setmanes):
            for dia_nom in self.dies:
                if not torns_per_dia_nom.get(dia_nom):
                    continue
                min_p, max_p = self.min_max_personal_per_dia(dia_nom)

                for nt, _, _ in torns_per_dia_nom[dia_nom]:
                    prob += pulp.lpSum([x[(t, w, dia_nom, nt)] for t in treballadors_actius]) >= min_p
                    prob += pulp.lpSum([x[(t, w, dia_nom, nt)] for t in treballadors_actius]) <= max_p

                    manca_llic[(w, dia_nom, nt)] = pulp.LpVariable(f"manca_llic_{w}_{dia_nom}_{nt}", lowBound=0)
                    prob += pulp.lpSum([x[(t, w, dia_nom, nt)] for t, info in treballadors_actius.items() if info['llicenciat']]) + manca_llic[(w, dia_nom, nt)] >= self.min_llicenciats_per_dia(dia_nom)
                    funcio_obj += PES_MANCA_LLICENCIAT * manca_llic[(w, dia_nom, nt)]

        # ------------------------------------------------------------------
        # ROTACIÓ CÍCLICA DE CAPS DE SETMANA (ÚNIC PARÀMETRE) + GARANTITS
        # ------------------------------------------------------------------
        dies_cs_oberts = [d for d in DIES_CAP_SETMANA if torns_per_dia_nom.get(d)]
        totals_rotacio = {}

        if dies_cs_oberts:
            for t, info in treballadors_actius.items():
                if self.te_patro_fix_cap_setmana(info):
                    continue

                caps_setmana_treballats = []
                for w in range(num_setmanes):
                    var_worked_weekend = pulp.LpVariable(f"weekend_{t}_w{w}", cat='Binary')
                    caps_setmana_treballats.append(var_worked_weekend)

                    torns_cap_setmana = [x[(t, w, d, nt)] for d in dies_cs_oberts for nt, _, _ in torns_per_dia_nom[d] if (t, w, d, nt) in x]
                    if torns_cap_setmana:
                        prob += pulp.lpSum(torns_cap_setmana) <= len(torns_cap_setmana) * var_worked_weekend
                        prob += var_worked_weekend <= pulp.lpSum(torns_cap_setmana)

                if caps_setmana_treballats:
                    total_t = pulp.lpSum(caps_setmana_treballats)
                    prob += total_t <= 1
                    # NOU: si el treballador té marcat "garantir ≥1 cap de setmana",
                    # es converteix en una restricció DURA: exactament 1 (no 0).
                    if info.get('garantir_cap_setmana', False):
                        prob += total_t >= 1
                    totals_rotacio[t] = total_t
                    funcio_obj += -PES_BONUS_ROTACIO * total_t

            noms_rotacio = list(totals_rotacio.keys())
            if len(noms_rotacio) >= 2:
                max_var = pulp.LpVariable("rot_max", lowBound=0)
                min_var = pulp.LpVariable("rot_min", lowBound=0)
                for t in noms_rotacio:
                    prob += max_var >= totals_rotacio[t]
                    prob += min_var <= totals_rotacio[t]
                prob += max_var - min_var <= DIFERENCIA_MAX_EQUITAT_ROTACIO

        prob += funcio_obj

        for t, info in treballadors_actius.items():
            patro = info.get('patro_semanal', {})
            for w in range(num_setmanes):
                for dia_nom in self.dies:
                    if torns_per_dia_nom.get(dia_nom):
                        val_especific = patro.get(dia_nom, "-")
                        torns_disponibles = [nt for nt, _, _ in torns_per_dia_nom[dia_nom]]

                        if val_especific == "Només Matí":
                            if "Tarda" in torns_disponibles:
                                prob += x[(t, w, dia_nom, "Tarda")] == 0
                        elif val_especific == "Només Tarda":
                            if "Matí" in torns_disponibles:
                                prob += x[(t, w, dia_nom, "Matí")] == 0
                        elif val_especific == "Tot el dia":
                            for nt, _, _ in torns_per_dia_nom[dia_nom]:
                                prob += x[(t, w, dia_nom, nt)] == 1
                        elif val_especific in ["No ve", "Festa"]:
                            for nt, _, _ in torns_per_dia_nom[dia_nom]:
                                prob += x[(t, w, dia_nom, nt)] == 0
                        elif val_especific == "Només aquest dia":
                            prob += pulp.lpSum([x[(t, w, dia_nom, nt)] for nt in torns_disponibles]) >= 1

        for t, info in treballadors_actius.items():
            permetre_extres = info.get('extres', False)
            target_cicle = target_cicle_dict.get(t, 0.0)

            hores_assignades_cicle = [
                x[(t, w, dia_nom, nt)] * durada
                for w in range(num_setmanes)
                for dia_nom in self.dies
                for nt, durada, _ in torns_per_dia_nom.get(dia_nom, [])
                if (t, w, dia_nom, nt) in x
            ]

            if hores_assignades_cicle:
                prob += pulp.lpSum(hores_assignades_cicle) - target_cicle == extra_h[t] - deficit_h[t]
                if not permetre_extres:
                    prob += extra_h[t] == 0

        num_fils = max(1, os.cpu_count() or 1)
        def _ruta_cbc():
            if getattr(sys, 'frozen', False):
                base = os.path.dirname(sys.executable)
            else:
                base = os.path.dirname(os.path.abspath(__file__))
            return os.path.join(base, "bin", "cbc.exe")

        solver = pulp.PULP_CBC_CMD(msg=False, timeLimit=40, threads=num_fils)
        prob.solve(solver)

        te_solucio_valida = any(
            pulp.value(var) is not None and pulp.value(var) > 0.5
            for var in x.values()
        ) if x else False

        if prob.status == pulp.LpStatusInfeasible:
            return None, (
                "❌ No s'han trobat solucions vàlides amb aquestes restriccions (problema infactible).\n\n"
                "Prova de relaxar algun paràmetre de la pestanya 2 (personal, rotació, garantits) o revisa els patrons setmanals."
            ), None, None

        if prob.status == pulp.LpStatusUndefined and not te_solucio_valida:
            return None, (
                "❌ El càlcul ha trigat massa i no s'ha pogut completar a temps.\n\n"
                "Prova d'escurçar el cicle o relaxar algun paràmetre."
            ), None, None

        self.dies_sense_llicenciat = set()
        for (w, dia_nom, nt), var in manca_llic.items():
            val = pulp.value(var)
            if val is not None and val > 0.5:
                self.dies_sense_llicenciat.add((w, dia_nom))

        res_quadrant = {}
        res_resum_dict = {}

        for t in treballadors_actius:
            h_totals = 0.0
            res_quadrant[t] = {}

            for w in range(num_setmanes):
                res_quadrant[t][w] = {}
                for dia_nom in self.dies:
                    torns_fets = []
                    for nt, durada, _ in torns_per_dia_nom.get(dia_nom, []):
                        if (t, w, dia_nom, nt) in x and pulp.value(x[(t, w, dia_nom, nt)]) is not None and pulp.value(x[(t, w, dia_nom, nt)]) > 0.5:
                            torns_fets.append(nt)
                            h_totals += durada

                    res_quadrant[t][w][dia_nom] = " + ".join(torns_fets) if torns_fets else "Festa"

            val_ext = pulp.value(extra_h[t])
            val_def = pulp.value(deficit_h[t])

            res_resum_dict[t] = {
                'fetes': h_totals,
                'extres': val_ext if val_ext else 0.0,
                'deficit': val_def if val_def else 0.0
            }

        return res_quadrant, res_resum_dict, torns_per_dia_nom, target_setmana_dict

    # ----------------------------------------------------------------------
    # EXPORTACIÓ
    # ----------------------------------------------------------------------
    def _text_curt_cel_la(self, val):
        mapping = {
            "Matí": "Matí",
            "Tarda": "Tarda",
            "Matí + Tarda": "M+T",
            "Festa": "-",
            "-": "-"
        }
        return mapping.get(val, val)

    def _construir_vista_torns_setmana(self, w):
        """Retorna (fila_mati, fila_tarda) per a la setmana w, en el mateix
        format que la vista 'Per Torn' de la pestanya de resultats."""
        fila_mati, fila_tarda = [], []
        for dia_nom in self.dies:
            torns_dia = self.torns_per_dia_nom.get(dia_nom, [])
            nts_disponibles = [nt for nt, _, _ in torns_dia]
            if not torns_dia:
                fila_mati.append("-")
                fila_tarda.append("-")
                continue
            pers_mati, pers_tarda = [], []
            for t, setmanes_dict in self.res_quadrant_cicle.items():
                val = setmanes_dict.get(w, {}).get(dia_nom, "Festa")
                if val in ["Matí", "Matí + Tarda"]:
                    pers_mati.append(t)
                if val in ["Tarda", "Matí + Tarda"]:
                    pers_tarda.append(t)
            fila_mati.append(", ".join(pers_mati) if pers_mati else "—") if "Matí" in nts_disponibles else fila_mati.append("-")
            fila_tarda.append(", ".join(pers_tarda) if pers_tarda else "—") if "Tarda" in nts_disponibles else fila_tarda.append("-")
        return fila_mati, fila_tarda

    def _hores_per_setmana_treballador(self, t):
        resultat = {}
        setmanes_dict = self.res_quadrant_cicle.get(t, {})
        for w, dies_dict in setmanes_dict.items():
            h = 0.0
            for dia_nom, val in dies_dict.items():
                torns_dia = self.torns_per_dia_nom.get(dia_nom, [])
                for nt, durada, _ in torns_dia:
                    if (val == "Matí" and nt == "Matí") or (val == "Tarda" and nt == "Tarda") or (val == "Matí + Tarda"):
                        h += durada
            resultat[w] = h
        return resultat    

    def exportar_pdf(self):
        if not self.res_quadrant_cicle:
            messagebox.showwarning("Atenció", "Primer has de generar un horari amb 'Generar Horari Cíclic'.")
            return
        if not REPORTLAB_OK:
            messagebox.showerror("Falta una llibreria", "Per exportar a PDF cal instal·lar 'reportlab'.\npip3 install reportlab")
            return

        ruta = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("Fitxer PDF", "*.pdf")],
            initialfile=f"Horari_Cicle_{self.num_setmanes_generat}setmanes.pdf"
        )
        if not ruta:
            return

        try:
            noms_treb = list(self.res_quadrant_cicle.keys())
            columnes = [(w, d) for w in range(self.num_setmanes_generat) for d in self.dies]

            colors_torn = {
                "Matí": colors.HexColor("#fef9c3"),
                "Tarda": colors.HexColor("#fed7aa"),
                "Matí + Tarda": colors.HexColor("#bfdbfe"),
                "Festa": colors.HexColor("#dcfce7"),
                "-": colors.whitesmoke
            }

            marge = 1 * cm
            mida_pagina = landscape(A3)
            amplada_pagina = mida_pagina[0]
            amplada_util = amplada_pagina - 2 * marge

            amplada_nom = 2.6 * cm
            num_cols = max(1, len(columnes))
            amplada_dia = (amplada_util - amplada_nom) / num_cols
            amplada_dia = min(amplada_dia, 2.2 * cm)
            amplada_dia = max(amplada_dia, 0.55 * cm)

            mida_font_taula = 6.5
            if amplada_dia < 0.85 * cm:
                mida_font_taula = 5.0
            elif amplada_dia < 1.1 * cm:
                mida_font_taula = 5.8

            doc = SimpleDocTemplate(ruta, pagesize=mida_pagina, leftMargin=marge, rightMargin=marge, topMargin=marge, bottomMargin=marge)
            estils = getSampleStyleSheet()
            elements = []
            elements.append(Paragraph(f"Horari Cíclic de la Farmàcia — {self.num_setmanes_generat} setmana(es) (es repeteix indefinidament)", estils['Title']))
            elements.append(Spacer(1, 0.3 * cm))

            # --- Taula 1: vista per treballador (com fins ara) ---
            capçalera = ["Treballador/a"] + [f"S{w+1}-{d[:2]}" for (w, d) in columnes]
            dades_taula = [capçalera]
            estil_cmds = [
                ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e293b")),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTSIZE', (0, 0), (-1, -1), mida_font_taula),
                ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 2),
                ('RIGHTPADDING', (0, 0), (-1, -1), 2),
            ]

            for r_idx, t in enumerate(noms_treb, start=1):
                fila = [t]
                for c_idx, (w, d) in enumerate(columnes, start=1):
                    val = self.res_quadrant_cicle[t].get(w, {}).get(d, "-")
                    text = self._text_curt_cel_la(val)
                    fila.append(text)
                    bg = colors_torn.get(val, colors.whitesmoke)
                    estil_cmds.append(('BACKGROUND', (c_idx, r_idx), (c_idx, r_idx), bg))
                if r_idx % 2 == 0:
                    estil_cmds.append(('BACKGROUND', (0, r_idx), (0, r_idx), colors.HexColor("#eef2f7")))
                dades_taula.append(fila)

            col_widths = [amplada_nom] + [amplada_dia] * num_cols
            taula = Table(dades_taula, colWidths=col_widths, repeatRows=1)
            taula.setStyle(TableStyle(estil_cmds))
            elements.append(Paragraph("Vista per Treballador/a", estils['Heading2']))
            elements.append(taula)
            elements.append(Spacer(1, 0.5 * cm))

            # --- Taula 2: vista per torns (Matí/Tarda), una graella per setmana ---
            elements.append(Paragraph("Vista per Torns (Matí / Tarda) — Farmàcia", estils['Heading2']))
            for w in range(self.num_setmanes_generat):
                fila_mati, fila_tarda = self._construir_vista_torns_setmana(w)
                cap_setmana_torn = ["Torn"] + self.dies
                dades_torn = [cap_setmana_torn, ["🌅 Matí"] + fila_mati, ["🌆 Tarda"] + fila_tarda]
                taula_torn = Table(dades_torn, colWidths=[2.2 * cm] + [3.2 * cm] * len(self.dies))
                taula_torn.setStyle(TableStyle([
                    ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e293b")),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                    ('FONTSIZE', (0, 0), (-1, -1), 7),
                    ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
                    ('BACKGROUND', (0, 2), (-1, 2), colors.HexColor("#eef2f7")),
                ]))
                elements.append(Paragraph(f"Setmana {w + 1} de {self.num_setmanes_generat}", estils['Heading3']))
                elements.append(taula_torn)
                elements.append(Spacer(1, 0.3 * cm))

            # --- Taula 3: balanç d'hores del cicle ---
            elements.append(Paragraph("Balanç d'hores del cicle (detall setmanal)", estils['Heading2']))
            resum_header = ["Treballador/a", "Contracte (h/s)"] + [f"Setmana {w+1}" for w in range(self.num_setmanes_generat)] + ["Total Cicle", "Extres (h)"]
            resum_dades = [resum_header]
            for t, dades in self.res_resum_cicle.items():
                h_per_setmana = self._hores_per_setmana_treballador(t)
                fila = [t, f"{self.treballadors[t]['hores']:.1f}"]
                for w in range(self.num_setmanes_generat):
                    fila.append(f"{h_per_setmana.get(w, 0.0):.1f}")
                fila.append(f"{dades['fetes']:.1f}")
                fila.append(f"{dades['extres']:.1f}")
                resum_dades.append(fila)
            taula_resum = Table(resum_dades, repeatRows=1)
            taula_resum.setStyle(TableStyle([
                ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e293b")),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTSIZE', (0, 0), (-1, -1), 7.5),
                ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
            ]))
            elements.append(taula_resum)

            doc.build(elements)
            messagebox.showinfo("Exportació completada", f"S'ha generat el PDF correctament a:\n{ruta}")
        except Exception as e:
            messagebox.showerror("Error d'exportació", f"No s'ha pogut generar el PDF:\n{str(e)}")

    def exportar_excel(self):
        if not self.res_quadrant_cicle:
            messagebox.showwarning("Atenció", "Primer has de generar un horari amb 'Generar Horari Cíclic'.")
            return
        if not OPENPYXL_OK:
            messagebox.showerror("Falta una llibreria", "Per exportar a Excel cal instal·lar 'openpyxl'.\npip3 install openpyxl")
            return

        ruta = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Fitxer Excel", "*.xlsx")],
            initialfile=f"Horari_Cicle_{self.num_setmanes_generat}setmanes.xlsx"
        )
        if not ruta:
            return

        try:
            noms_treb = list(self.res_quadrant_cicle.keys())
            columnes = [(w, d) for w in range(self.num_setmanes_generat) for d in self.dies]

            fills_torn = {
                "Matí": PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),
                "Tarda": PatternFill(start_color="FED7AA", end_color="FED7AA", fill_type="solid"),
                "Matí + Tarda": PatternFill(start_color="BFDBFE", end_color="BFDBFE", fill_type="solid"),
                "Festa": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
                "-": PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid"),
            }
            fill_zebra = PatternFill(start_color="EEF2F7", end_color="EEF2F7", fill_type="solid")

            fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            font_header = Font(color="FFFFFF", bold=True)
            thin = Side(style="thin", color="CCCCCC")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            wb = Workbook()
            ws = wb.active
            ws.title = "Per Treballador"

            ws.cell(row=1, column=1, value=f"Horari Cíclic de la Farmàcia — {self.num_setmanes_generat} setmana(es) (es repeteix indefinidament)").font = Font(bold=True, size=14)
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columnes) + 1)

            header_row = 3
            ws.cell(row=header_row, column=1, value="Treballador/a").fill = fill_header
            ws.cell(row=header_row, column=1).font = font_header
            ws.cell(row=header_row, column=1).border = border
            for c_idx, (w, d) in enumerate(columnes, start=2):
                cell = ws.cell(row=header_row, column=c_idx, value=f"S{w+1}-{d[:2]}")
                cell.fill = fill_header
                cell.font = font_header
                cell.alignment = Alignment(horizontal="center")
                cell.border = border

            for r_offset, t in enumerate(noms_treb):
                row = header_row + 1 + r_offset
                cell_nom = ws.cell(row=row, column=1, value=t)
                cell_nom.font = Font(bold=True)
                cell_nom.border = border
                if r_offset % 2 == 1:
                    cell_nom.fill = fill_zebra
                for c_offset, (w, d) in enumerate(columnes):
                    col = 2 + c_offset
                    val = self.res_quadrant_cicle[t].get(w, {}).get(d, "-")
                    text = self._text_curt_cel_la(val)
                    cell = ws.cell(row=row, column=col, value=text)
                    cell.alignment = Alignment(horizontal="center")
                    cell.border = border
                    cell.fill = fills_torn.get(val, fills_torn["-"])

            ws.column_dimensions['A'].width = 16
            for c_idx in range(2, len(columnes) + 2):
                ws.column_dimensions[get_column_letter(c_idx)].width = 7
            ws.freeze_panes = ws.cell(row=header_row + 1, column=2)

            # --- Fulla 2: vista per torns (Matí/Tarda), una graella per setmana ---
            ws_torn = wb.create_sheet("Per Torn (Farmàcia)")
            r = 1
            ws_torn.cell(row=r, column=1, value="Vista per Torns (Matí / Tarda) — Farmàcia").font = Font(bold=True, size=14)
            r += 2
            for w in range(self.num_setmanes_generat):
                fila_mati, fila_tarda = self._construir_vista_torns_setmana(w)
                ws_torn.cell(row=r, column=1, value=f"Setmana {w + 1} de {self.num_setmanes_generat}").font = Font(bold=True, size=11)
                r += 1
                ws_torn.cell(row=r, column=1, value="Torn").fill = fill_header
                ws_torn.cell(row=r, column=1).font = font_header
                ws_torn.cell(row=r, column=1).border = border
                for c_idx, dia_nom in enumerate(self.dies, start=2):
                    cell = ws_torn.cell(row=r, column=c_idx, value=dia_nom)
                    cell.fill = fill_header
                    cell.font = font_header
                    cell.border = border
                    cell.alignment = Alignment(horizontal="center")
                r += 1
                for etiqueta, fila in (("🌅 Matí", fila_mati), ("🌆 Tarda", fila_tarda)):
                    ws_torn.cell(row=r, column=1, value=etiqueta).border = border
                    ws_torn.cell(row=r, column=1).font = Font(bold=True)
                    for c_idx, val in enumerate(fila, start=2):
                        cell = ws_torn.cell(row=r, column=c_idx, value=val)
                        cell.border = border
                        cell.alignment = Alignment(horizontal="center")
                    r += 1
                r += 1
            ws_torn.column_dimensions['A'].width = 10
            for c_idx in range(2, len(self.dies) + 2):
                ws_torn.column_dimensions[get_column_letter(c_idx)].width = 22

            # --- Fulla 3: resum del cicle ---
            ws2 = wb.create_sheet("Resum Cicle")
            headers2 = ["Treballador/a", "Contracte (h/s)"] + [f"Setmana {w+1}" for w in range(self.num_setmanes_generat)] + ["Total Cicle", "Extres (h)"]
            for c_idx, h in enumerate(headers2, start=1):
                cell = ws2.cell(row=1, column=c_idx, value=h)
                cell.fill = fill_header
                cell.font = font_header
                cell.border = border
                cell.alignment = Alignment(horizontal="center")

            for r_offset, (t, dades) in enumerate(self.res_resum_cicle.items(), start=2):
                fila_fill = fill_zebra if r_offset % 2 == 0 else None
                h_per_setmana = self._hores_per_setmana_treballador(t)
                valors = [t, round(self.treballadors[t]['hores'], 1)]
                for w in range(self.num_setmanes_generat):
                    valors.append(round(h_per_setmana.get(w, 0.0), 1))
                valors.append(round(dades['fetes'], 1))
                valors.append(round(dades['extres'], 1))
                for c_idx, v in enumerate(valors, start=1):
                    c = ws2.cell(row=r_offset, column=c_idx, value=v)
                    c.border = border
                    if fila_fill:
                        c.fill = fila_fill

            for c_idx in range(1, len(headers2) + 1):
                ws2.column_dimensions[get_column_letter(c_idx)].width = 16

            wb.save(ruta)
            messagebox.showinfo("Exportació completada", f"S'ha generat l'Excel correctament a:\n{ruta}")
        except Exception as e:
            messagebox.showerror("Error d'exportació", f"No s'ha pogut generar l'Excel:\n{str(e)}")

    # ----------------------------------------------------------------------
    # PERSISTÈNCIA
    # ----------------------------------------------------------------------
    def guardar_dades(self):
        dades = {
            'treballadors': self.treballadors,
            'config': {
                'min_personal_setmana': self.min_personal_setmana_var.get(),
                'max_personal_setmana': self.max_personal_setmana_var.get(),
                'min_personal_cap_setmana': self.min_personal_cap_setmana_var.get(),
                'max_personal_cap_setmana': self.max_personal_cap_setmana_var.get(),
                'min_llicenciats': self.min_llicenciats_var.get(),
                'min_llicenciats_cap_setmana': self.min_llicenciats_cap_setmana_var.get(),
                'num_setmanes_cicle': self.num_setmanes_cicle_var.get(),
                'inici_canvi': self.inici_canvi_var.get(),
                'fi_canvi': self.fi_canvi_var.get(),
                'durada_min_torn': self.durada_min_torn_var.get(),
                'durada_max_torn': self.durada_max_torn_var.get(),
                'mida_font': self.mida_font.get()
            },
            'horaris': {dia: {'inici': self.horaris_ui[dia]['inici'].get(), 'fi': self.horaris_ui[dia]['fi'].get()} for dia in self.dies}
        }
        with open(FITXER_DADES, 'w', encoding='utf-8') as f:
            json.dump(dades, f, ensure_ascii=False, indent=4)

    def carregar_dades(self):
        if not os.path.exists(FITXER_DADES):
            return
        try:
            with open(FITXER_DADES, 'r', encoding='utf-8') as f:
                dades = json.load(f)

            self.treballadors = dades.get('treballadors', {})

            cfg = dades.get('config', {})

            min_p_legacy = cfg.get('min_personal', 2)
            max_p_legacy = cfg.get('max_personal', 4)

            self.min_personal_setmana_var.set(cfg.get('min_personal_setmana', min_p_legacy))
            self.max_personal_setmana_var.set(cfg.get('max_personal_setmana', max_p_legacy))
            self.min_personal_cap_setmana_var.set(cfg.get('min_personal_cap_setmana', max(1, min_p_legacy - 1)))
            self.max_personal_cap_setmana_var.set(cfg.get('max_personal_cap_setmana', max(1, max_p_legacy - 1)))

            self.min_llicenciats_var.set(cfg.get('min_llicenciats', 1))
            self.min_llicenciats_cap_setmana_var.set(cfg.get('min_llicenciats_cap_setmana', 1))
            self.num_setmanes_cicle_var.set(cfg.get('num_setmanes_cicle', cfg.get('max_caps_setmana', 2)))
            self.inici_canvi_var.set(cfg.get('inici_canvi', "13:30"))
            self.fi_canvi_var.set(cfg.get('fi_canvi', "14:30"))
            self.durada_min_torn_var.set(cfg.get('durada_min_torn', 4.0))
            self.durada_max_torn_var.set(cfg.get('durada_max_torn', 8.0))
            self.mida_font.set(cfg.get('mida_font', 10))

            horaris = dades.get('horaris', {})
            for dia, h in horaris.items():
                if dia in self.horaris_ui:
                    self.horaris_ui[dia]['inici'].set(h.get('inici', ''))
                    self.horaris_ui[dia]['fi'].set(h.get('fi', ''))

            self.actualitzar_llista_treballadors()
            self.aplicar_estils_font()
            self.actualitzar_info_rotacio()
        except Exception as e:
            print("Error en carregar les dades:", e)

    def on_closing(self):
        self.guardar_dades()
        self.root.destroy()


if __name__ == "__main__":
    root = tk.Tk()
    app = AppFarmaciaCiclica(root)
    root.mainloop()
